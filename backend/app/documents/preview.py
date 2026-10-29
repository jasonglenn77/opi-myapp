"""
In-app document preview (viewer modal): turn a stored file into something the
browser can show without downloading it.

  pdf / images      -> served inline by /api/documents/file/{id}/content
  HEIC / HEIF       -> converted to JPEG on the way out (iPhone photos)
  xlsx / xlsm       -> HTML tables (every visible sheet, or one named sheet)
  docx              -> HTML (mammoth)
  eml               -> headers + the message body
  csv / txt         -> table / plain text
Everything HTML is rendered by the viewer in a sandboxed iframe (no scripts).
"""
import csv
import email
import html
import io
from datetime import date, datetime, time
from email import policy

IMAGE_EXT = {"jpg", "jpeg", "png", "gif", "webp", "bmp"}
HEIC_EXT = {"heic", "heif"}
VIDEO_EXT = {"mp4", "mov", "webm", "m4v"}
HTML_EXT = {"xlsx", "xlsm", "docx", "eml", "csv", "txt"}
MIME = {"pdf": "application/pdf", "jpg": "image/jpeg", "jpeg": "image/jpeg", "png": "image/png",
        "gif": "image/gif", "webp": "image/webp", "bmp": "image/bmp", "mp4": "video/mp4",
        "mov": "video/quicktime", "webm": "video/webm", "m4v": "video/mp4"}
MAX_ROWS, MAX_COLS = 600, 40

_CSS = """
body{font:13px/1.45 -apple-system,Segoe UI,Roboto,Helvetica,Arial,sans-serif;color:#1c2430;margin:0;padding:16px;background:#fff}
h2{font-size:13px;margin:18px 0 6px;color:#3b4656;text-transform:uppercase;letter-spacing:.04em}
h2:first-child{margin-top:0}
.wrap{overflow-x:auto;margin-bottom:8px}
table.x{border-collapse:collapse;font-size:12px;font-variant-numeric:tabular-nums}
table.x td{border:1px solid #dfe3e8;padding:3px 6px;white-space:nowrap;vertical-align:top;max-width:420px;overflow:hidden;text-overflow:ellipsis}
table.x td.n{text-align:right}
table.x tr.rh td{background:#f4f6f8;color:#8a94a3;font-size:10px;text-align:center}
table.x td.rn{background:#f4f6f8;color:#8a94a3;font-size:10px;text-align:right}
.note{color:#7a8594;font-size:11px;margin:4px 0 0}
.mail dt{float:left;width:64px;color:#7a8594;font-weight:600}
.mail dd{margin:0 0 4px 72px}
.mail{border-bottom:1px solid #e5e8ec;padding-bottom:10px;margin-bottom:12px}
pre{white-space:pre-wrap;font:12px/1.5 ui-monospace,Consolas,monospace}
@media print{body{padding:0}.wrap{overflow:visible}}
"""


def ext_of(name):
    n = (name or "").lower().rsplit(".", 1)
    return n[1] if len(n) == 2 else ""


def kind_of(name, content_type=None):
    e = ext_of(name)
    ct = (content_type or "").lower()
    if e == "pdf" or ct == "application/pdf":
        return "pdf"
    if e in IMAGE_EXT or e in HEIC_EXT or ct.startswith("image/"):
        return "image"
    if e in VIDEO_EXT or ct.startswith("video/"):
        return "video"
    if e in HTML_EXT:
        return "html"
    return "none"


def _page(title, body):
    return (f"<!doctype html><html><head><meta charset='utf-8'><title>{html.escape(title or '')}</title>"
            f"<style>{_CSS}</style></head><body>{body}</body></html>")


def heic_to_jpeg(data):
    from PIL import Image
    import pillow_heif
    pillow_heif.register_heif_opener()
    im = Image.open(io.BytesIO(data))
    im = im.convert("RGB")
    out = io.BytesIO()
    im.save(out, "JPEG", quality=85)
    return out.getvalue()


# --- spreadsheets -----------------------------------------------------------
def _fmt(v, number_format):
    if v is None:
        return "", False
    nf = (number_format or "").lower()
    if isinstance(v, datetime):
        return (v.strftime("%-m/%-d/%Y") if v.time() == time(0) else v.strftime("%-m/%-d/%Y %-I:%M %p")), True
    if isinstance(v, date):
        return v.strftime("%-m/%-d/%Y"), True
    if isinstance(v, bool):
        return ("TRUE" if v else "FALSE"), False
    if isinstance(v, (int, float)):
        if "%" in nf:
            return f"{v * 100:,.{1 if '0.0' in nf else 0}f}%", True
        if "$" in nf or "[$" in nf or "accounting" in nf:
            s = f"${abs(v):,.2f}" if ".00" in nf or "0.00" in nf else f"${abs(v):,.0f}"
            return (f"({s})" if v < 0 and "(" in nf else ("-" + s if v < 0 else s)), True
        if "0.00" in nf:
            return f"{v:,.2f}", True
        if "#,##0" in nf:
            return f"{v:,.0f}", True
        return (f"{v:,.2f}".rstrip("0").rstrip(".") if isinstance(v, float) else str(v)), True
    return str(v), False


def _color(c):
    try:
        if c is not None and c.type == "rgb" and isinstance(c.rgb, str) and len(c.rgb) == 8 and c.rgb not in ("00000000", "FF000000"):
            return "#" + c.rgb[2:]
    except Exception:
        pass
    return None


def _sheet_html(ws):
    max_r, max_c = min(ws.max_row or 0, MAX_ROWS), min(ws.max_column or 0, MAX_COLS)
    # trim trailing empty rows / columns
    last_r = last_c = 0
    for row in ws.iter_rows(min_row=1, max_row=max_r, max_col=max_c):
        for cell in row:
            if cell.value not in (None, ""):
                last_r = max(last_r, cell.row)
                last_c = max(last_c, cell.column)
    if not last_r:
        return "<p class='note'>(empty sheet)</p>"
    span, skip = {}, set()
    for rng in ws.merged_cells.ranges:
        if rng.min_row > last_r or rng.min_col > last_c:
            continue
        span[(rng.min_row, rng.min_col)] = (min(rng.max_row, last_r) - rng.min_row + 1, min(rng.max_col, last_c) - rng.min_col + 1)
        for r in range(rng.min_row, min(rng.max_row, last_r) + 1):
            for c in range(rng.min_col, min(rng.max_col, last_c) + 1):
                if (r, c) != (rng.min_row, rng.min_col):
                    skip.add((r, c))
    from openpyxl.utils import get_column_letter
    hidden_cols = {c for c in range(1, last_c + 1)
                   if ws.column_dimensions[get_column_letter(c)].hidden}
    cols = [c for c in range(1, last_c + 1) if c not in hidden_cols]
    out = ["<div class='wrap'><table class='x'><tr class='rh'><td></td>"]
    out += [f"<td>{get_column_letter(c)}</td>" for c in cols]
    out.append("</tr>")
    for r in range(1, last_r + 1):
        if ws.row_dimensions[r].hidden:
            continue
        out.append(f"<tr><td class='rn'>{r}</td>")
        for c in cols:
            if (r, c) in skip:
                continue
            cell = ws.cell(r, c)
            txt, num = _fmt(cell.value, cell.number_format)
            st = []
            try:
                if cell.font is not None and cell.font.b:
                    st.append("font-weight:700")
                fc = _color(cell.font.color) if cell.font is not None else None
                if fc:
                    st.append(f"color:{fc}")
                if cell.fill is not None and cell.fill.fill_type == "solid":
                    bg = _color(cell.fill.fgColor)
                    if bg and bg.lower() != "#ffffff":
                        st.append(f"background:{bg}")
                al = cell.alignment.horizontal if cell.alignment is not None else None
                if al in ("center", "right", "left"):
                    st.append(f"text-align:{al}")
            except Exception:
                pass
            rs, cs = span.get((r, c), (1, 1))
            attrs = (f" rowspan='{rs}'" if rs > 1 else "") + (f" colspan='{cs}'" if cs > 1 else "")
            attrs += " class='n'" if num else ""
            attrs += f" style='{';'.join(st)}'" if st else ""
            out.append(f"<td{attrs} title='{html.escape(txt, quote=True)}'>{html.escape(txt)}</td>")
        out.append("</tr>")
    out.append("</table></div>")
    if (ws.max_row or 0) > MAX_ROWS or (ws.max_column or 0) > MAX_COLS:
        out.append(f"<p class='note'>Showing the first {MAX_ROWS} rows and {MAX_COLS} columns. Download the file for the rest.</p>")
    return "".join(out)


def xlsx_html(data, title, sheet=None):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    sheets = [wb[sheet]] if sheet and sheet in wb.sheetnames else [ws for ws in wb.worksheets if ws.sheet_state == "visible"]
    body = []
    for ws in sheets[:12]:
        if len(sheets) > 1:
            body.append(f"<h2>{html.escape(ws.title)}</h2>")
        body.append(_sheet_html(ws))
    if len(sheets) > 12:
        body.append(f"<p class='note'>{len(sheets) - 12} more sheets not shown. Download the file to see them.</p>")
    return _page(title, "".join(body))


def docx_html(data, title):
    import mammoth
    res = mammoth.convert_to_html(io.BytesIO(data))
    return _page(title, f"<div style='max-width:820px'>{res.value}</div>")


def eml_html(data, title):
    msg = email.message_from_bytes(data, policy=policy.default)
    head = "".join(f"<dt>{k}</dt><dd>{html.escape(str(msg.get(k) or ''))}</dd>"
                   for k in ("From", "To", "Cc", "Date", "Subject") if msg.get(k))
    body_part = msg.get_body(preferencelist=("html", "plain"))
    body = ""
    if body_part is not None:
        content = body_part.get_content()
        body = content if body_part.get_content_type() == "text/html" else f"<pre>{html.escape(content)}</pre>"
    atts = [p.get_filename() for p in msg.iter_attachments() if p.get_filename()]
    att = f"<p class='note'>Attachments: {html.escape(', '.join(atts))}</p>" if atts else ""
    return _page(title, f"<dl class='mail'>{head}</dl>{att}{body}")


def text_html(data, title, is_csv):
    s = data.decode("utf-8", errors="replace")
    if not is_csv:
        return _page(title, f"<pre>{html.escape(s)}</pre>")
    rows = list(csv.reader(io.StringIO(s)))[:MAX_ROWS]
    t = "".join("<tr>" + "".join(f"<td>{html.escape(c)}</td>" for c in r[:MAX_COLS]) + "</tr>" for r in rows)
    return _page(title, f"<div class='wrap'><table class='x'>{t}</table></div>")


def render_html(name, data, sheet=None):
    e = ext_of(name)
    if e in ("xlsx", "xlsm"):
        return xlsx_html(data, name, sheet)
    if e == "docx":
        return docx_html(data, name)
    if e == "eml":
        return eml_html(data, name)
    return text_html(data, name, e == "csv")
