"""
Ashley's Tracker — the office's estimates & progress-invoicing summary and the
per-PM project schedule, replacing the "Progress Invoice" workbook.

Invoicing view: one row per project with its status, PMs, dates, accepted
estimates and their invoice schedules. Billing numbers come from the SAME
composer the Billing & Schedule tab uses (billing.routes._compose_estimates),
so both screens always agree; edits go through /api/invoices (shared rows,
audited). Composing every project takes seconds, so each project's result is
cached in-process and recomputed only when its billing signature changes.

PM schedule view: every active PM assignment with its schedule window, for the
per-PM table and timeline.
"""
import json
import threading
from datetime import date, timedelta

from fastapi import APIRouter, Depends, HTTPException
from sqlalchemy import bindparam, text

from app.db import engine
from app.auth import get_current_user
from app.permissions import has_capability, PAGE_CUSTOMERS
from app.billing.routes import (
    _compose_estimates, _ensure_estimate_billing, _apply_assignment_default_crew,
    _ensure_invoice_schedules, _default_crew_for_project, _project_meta,
)
from app.invoices.routes import _project_ctx
from app.billing.align import _unbilled_balance

router = APIRouter(prefix="/api/tracker", tags=["tracker"])

PAST_STATUSES = ("completed", "canceled")
SOON_DAYS = 14

_cache = {}            # entity_id -> (signature, composed billing dict)
_cache_lock = threading.Lock()


def _require(user):
    if not has_capability(user, PAGE_CUSTOMERS):
        raise HTTPException(status_code=403, detail="Insufficient permissions")


def _d(v):
    return str(v)[:10] if v else None


def _initials(first, last):
    return ((first or "")[:1] + (last or "")[:1]).upper() or "?"


# ── bulk lookups ────────────────────────────────────────────────────────────
def _projects(conn, only=None):
    where = "AND qc.qbo_id = :only" if only else ""
    rows = conn.execute(text(f"""
        SELECT qc.qbo_id, qc.id AS qc_id, qc.display_name, pc.display_name AS customer
        FROM qbo_customers qc LEFT JOIN qbo_customers pc ON pc.qbo_id = qc.parent_qbo_id
        WHERE qc.is_project = 1 {where}
    """), {"only": only} if only else {}).mappings().all()
    return {str(r["qbo_id"]): dict(r) for r in rows}


def _schedule(conn, only=None):
    where = "WHERE qc.qbo_id = :only" if only else ""
    rows = conn.execute(text(f"""
        WITH ranked AS (
          SELECT qc.qbo_id, psi.status, psi.start_date, psi.end_date,
                 DATE_ADD(psi.end_date, INTERVAL COALESCE(psi.overage_days, 0) DAY) AS true_end,
                 ROW_NUMBER() OVER (PARTITION BY qc.qbo_id
                   ORDER BY psi.start_date IS NULL, psi.start_date DESC, psi.sort_order DESC, psi.id DESC) AS rn
          FROM project_schedule_items psi
          JOIN projects p ON p.id = psi.project_id
          JOIN qbo_customers qc ON qc.id = p.qbo_customer_id
          {where}
        )
        SELECT qbo_id, MAX(CASE WHEN rn = 1 THEN status END) AS status,
               MIN(start_date) AS start_date, MAX(end_date) AS end_date, MAX(true_end) AS true_end
        FROM ranked GROUP BY qbo_id
    """), {"only": only} if only else {}).mappings().all()
    return {str(r["qbo_id"]): r for r in rows}


def _pms(conn, only=None):
    where = "AND qc.qbo_id = :only" if only else ""
    rows = conn.execute(text(f"""
        SELECT qc.qbo_id, m.id, m.first_name, m.last_name, MAX(spm.is_primary) AS is_primary
        FROM project_schedule_item_project_managers spm
        JOIN project_schedule_items psi ON psi.id = spm.schedule_item_id
        JOIN projects p ON p.id = psi.project_id
        JOIN qbo_customers qc ON qc.id = p.qbo_customer_id
        JOIN project_managers m ON m.id = spm.project_manager_id
        WHERE spm.unassigned_at IS NULL {where}
        GROUP BY qc.qbo_id, m.id, m.first_name, m.last_name
        ORDER BY is_primary DESC, m.first_name
    """), {"only": only} if only else {}).mappings().all()
    out = {}
    for r in rows:
        out.setdefault(str(r["qbo_id"]), []).append({
            "id": r["id"], "name": f"{r['first_name'] or ''} {r['last_name'] or ''}".strip(),
            "initials": _initials(r["first_name"], r["last_name"]), "primary": bool(r["is_primary"])})
    return out


def _estimate_dates(conn, only=None):
    where = "AND customer_qbo_id = :only" if only else ""
    rows = conn.execute(text(f"""
        SELECT qbo_id, MAX(txn_date) AS txn_date FROM qbo_transactions
        WHERE entity_type = 'Estimate' {where} GROUP BY qbo_id
    """), {"only": only} if only else {}).all()
    return {str(r[0]): r[1] for r in rows}


def _signatures(conn, only=None):
    """Cheap per-project fingerprint of everything _compose_estimates reads."""
    p = {"only": only} if only else {}
    sig = {}

    def add(rows):
        for r in rows:
            sig.setdefault(str(r[0]), []).append(tuple(str(x) for x in r[1:]))

    add(conn.execute(text(f"""
        SELECT s.entity_id, COUNT(m.id), MAX(m.updated_at), COALESCE(SUM(m.amount), 0), COUNT(DISTINCT s.id)
        FROM project_invoice_schedules s LEFT JOIN project_invoice_milestones m ON m.schedule_id = s.id
        {"WHERE s.entity_id = :only" if only else ""} GROUP BY s.entity_id
    """), p).all())
    add(conn.execute(text(f"""
        SELECT customer_qbo_id, MAX(id), COUNT(*) FROM qbo_transactions
        WHERE entity_type IN ('Invoice', 'Estimate') {"AND customer_qbo_id = :only" if only else ""}
        GROUP BY customer_qbo_id
    """), p).all())
    add(conn.execute(text(f"""
        SELECT entity_id, MAX(updated_at), COUNT(*) FROM project_estimate_billing
        {"WHERE entity_id = :only" if only else ""} GROUP BY entity_id
    """), p).all())
    try:   # change orders (pending COs on the billing composition)
        add(conn.execute(text(f"""
            SELECT project_qbo_id, MAX(updated_at), COUNT(*) FROM project_change_orders
            {"WHERE project_qbo_id = :only" if only else ""} GROUP BY project_qbo_id
        """), p).all())
    except Exception:
        conn.rollback()
    try:   # 0079 expected-paid overrides on actual invoices
        add(conn.execute(text(f"""
            SELECT entity_id, MAX(updated_at), COUNT(*), COUNT(expected_paid_date) FROM qbo_invoice_expected_paid
            {"WHERE entity_id = :only" if only else ""} GROUP BY entity_id
        """), p).all())
    except Exception:
        pass
    return {k: tuple(sorted(v)) for k, v in sig.items()}


def _needs_setup(conn, only=None):
    """Active/upcoming projects with an accepted estimate that has no invoice
    schedule yet — the same gap Billing & Schedule fills when it's opened."""
    rows = conn.execute(text(f"""
        SELECT DISTINCT t.customer_qbo_id
        FROM qbo_transactions t
        JOIN qbo_customers qc ON qc.qbo_id = t.customer_qbo_id AND qc.is_project = 1
        WHERE t.entity_type = 'Estimate' AND t.total_amt > 0
          AND JSON_UNQUOTE(JSON_EXTRACT(t.raw_json, '$.TxnStatus')) IN ('Accepted', 'Converted')
          AND NOT EXISTS (SELECT 1 FROM project_invoice_schedules s
                          WHERE s.entity_id = t.customer_qbo_id AND s.estimate_qbo_id = t.qbo_id)
          {"AND t.customer_qbo_id = :only" if only else ""}
    """), {"only": only} if only else {}).all()
    return [str(r[0]) for r in rows]


def _setup(entity_id):
    """Run Billing & Schedule's create-if-missing steps (billing headers + one
    invoice schedule per accepted estimate) so both screens show the same plan."""
    with engine.begin() as conn:
        ctx = _project_ctx(conn, entity_id)
        if not ctx:
            return
        meta = _project_meta(conn, entity_id) or {}
        _ensure_estimate_billing(conn, entity_id, _default_crew_for_project(conn, entity_id, meta))
        _apply_assignment_default_crew(conn, entity_id)
        _ensure_invoice_schedules(conn, entity_id, ctx)


_preloaded = {"done": False}


def _preload():
    """0089: a fresh worker (after a deploy / restart) starts from every
    project's last saved composition in one query, instead of recomposing all
    ~600 projects (~23s); _billing still recomposes any whose signature moved."""
    if _preloaded["done"]:
        return
    try:
        from app import snapshots
        snaps = snapshots.get_many("tracker:")
        with _cache_lock:
            for k, (payload, sig) in snaps.items():
                _cache.setdefault(k.split(":", 1)[1], (sig, payload))
    except Exception:
        pass
    _preloaded["done"] = True


def _billing(conn, entity_id, sig):
    sj = json.dumps(sig)
    with _cache_lock:
        hit = _cache.get(entity_id)
    if hit and hit[0] == sj:
        return hit[1]
    comp = _compose_estimates(conn, entity_id)
    with _cache_lock:
        _cache[entity_id] = (sj, comp)
    try:   # shared with the other workers and kept across restarts
        from app import snapshots
        snapshots.put(f"tracker:{entity_id}", comp, sj)
    except Exception:
        pass
    return comp


def refresh_snapshot():
    """Nightly warm-up: recompose (and save) every project whose data changed."""
    return len(_build() or [])


# ── row composition ─────────────────────────────────────────────────────────
def _row(eid, proj, sched, pms, est_dates, comp, today):
    status = (sched["status"] if sched else None) or "needs_attention"
    start = _d(sched["start_date"]) if sched else None
    end = _d(sched["true_end"] or sched["end_date"]) if sched else None
    days = None
    if start and end:
        days = (date.fromisoformat(end) - date.fromisoformat(start)).days + 1

    estimates, milestones = [], []
    for a in comp.get("accepted", []):
        ms = []
        for m in a.get("milestones", []):
            # A partly-sent invoice whose leftover is rounding-sized counts as sent.
            left = float(m.get("remaining") or 0)
            if m["tier"] == "partial" and left < max(100.0, 0.01 * float(m.get("amount") or 0)):
                m = {**m, "tier": "committed", "status_label": "Sent · A/R"}
            unbilled = m["tier"] in ("scheduled", "partial")
            inv = m.get("invoice_date")
            flag = None
            if unbilled and inv and inv < today.isoformat():
                flag = "overdue_send"
            elif unbilled and inv and inv <= (today + timedelta(days=SOON_DAYS)).isoformat():
                flag = "send_soon"
            elif m["tier"] in ("committed", "partial") and m.get("due_date") and m["due_date"] < today.isoformat():
                flag = "payment_late"
            if m["tier"] == "retainage":
                flag = "retainage"   # held back until release — not a late payment
            mm = {k: m.get(k) for k in ("id", "seq", "label", "pct", "amount", "invoice_date", "due_date",
                                         "expected_paid_date", "tier", "status_label", "edited", "remaining",
                                         "retainage_held")}
            mm["flag"] = flag
            mm["estimate_doc"] = a.get("doc")
            ms.append(mm)
            milestones.append(mm)
        acts = a.get("invoice_actuals") or []
        e_inv = round(sum(float(x.get("amount") or 0) for x in acts), 2)
        e_open = round(sum(float(x.get("balance") or 0) for x in acts), 2)
        # partly billed with the rest due / undated / on a completed job (same
        # rule as Data Review's 'unbilled_balance')
        ub = _unbilled_balance(a.get("milestones") or [], acts, status, today)
        estimates.append({
            "qbo_id": a["qbo_id"], "doc": a.get("doc"), "date": _d(est_dates.get(str(a["qbo_id"]))),
            "value": a.get("value") or 0, "schedule_id": a.get("schedule_id"),
            "scheduled": a.get("invoice_subtotal") or 0,
            "invoiced": e_inv, "paid": round(e_inv - e_open, 2), "open_ar": e_open,
            "to_bill": round(max(0.0, (a.get("value") or 0) - e_inv), 2),
            "retainage_held": round(float(a.get("retainage_held") or 0), 2),
            "unbilled": ({"amount": ub["balance"], "date": ub["date"], "why": ub["why"]} if ub else None),
            "milestones": ms,
            "invoices": [{
                "doc": x.get("doc_number"), "date": x.get("txn_date"), "due": x.get("due_date"),
                "amount": x.get("amount"), "balance": x.get("balance"), "status": x.get("status"),
                "late": bool(float(x.get("balance") or 0) > 0.01 and x.get("due_date") and not x.get("retainage")
                             and str(x.get("due_date")) < today.isoformat()),
                "retainage": x.get("retainage") or 0.0,
                "estimate_doc": x.get("estimate_doc"),
                # 0079: expected-paid date on the actual invoice (Cash Flow uses it)
                "qbo_id": x.get("qbo_id"),
                "expected": x.get("expected_paid_date"),
                "expected_override": x.get("expected_override"),
            } for x in acts],
        })

    contract = round(sum(e["value"] for e in estimates), 2)
    invoiced = round(float(comp.get("invoiced_qbo") or 0), 2)
    open_ar = round(float(comp.get("open_ar") or 0), 2)
    collected = round(float(comp.get("collected") or 0), 2)
    remaining = round(max(0.0, contract - invoiced), 2)
    flags = sorted({m["flag"] for m in milestones if m["flag"]})
    retainage = round(sum(e["retainage_held"] for e in estimates), 2)
    if retainage > 0.5 and "retainage" not in flags:
        flags.append("retainage")
    ubs = [e["unbilled"] for e in estimates if e["unbilled"]]
    ub_row = ({"amount": round(sum(u["amount"] for u in ubs), 2),
                 "date": min((u["date"] for u in ubs if u["date"]), default=None),
                 "why": ("past" if any(u["why"] == "past" for u in ubs) else ubs[0]["why"])} if ubs else None)
    if ub_row:
        flags.append("unbilled_balance")
    pending = comp.get("pending") or []
    if pending:
        flags.append("pending_estimate")
    unbilled = sorted((m for m in milestones if m["tier"] in ("scheduled", "partial") and m.get("invoice_date")),
                      key=lambda m: m["invoice_date"])
    nxt = unbilled[0] if unbilled else None
    past = status in PAST_STATUSES
    relevant = (not past) or open_ar > 0.5 or bool(
        {"overdue_send", "send_soon", "payment_late", "unbilled_balance"} & set(flags))
    name = proj["display_name"] or ""
    num = name.split(" - ")[0].strip() if " - " in name else ""
    return {
        "entity_id": eid, "qc_id": proj.get("qc_id"), "name": name, "number": num, "customer": proj.get("customer"),
        "status": status, "past": past, "relevant": relevant,
        "pms": pms, "start": start, "end": end, "days": days,
        "contract": contract, "invoiced": invoiced, "collected": collected, "open_ar": open_ar,
        "remaining": remaining, "pct_invoiced": round(invoiced / contract * 100, 1) if contract else None,
        "retainage_held": retainage, "unbilled": ub_row,
        "estimates": estimates, "pending_estimates": pending,
        "next": ({"label": nxt["label"], "date": nxt["invoice_date"], "amount": nxt.get("remaining") or nxt["amount"],
                  "flag": nxt["flag"]} if nxt else None),
        "flags": flags,
    }


def _build(only=None):
    today = date.today()
    with engine.connect() as conn:
        projs = _projects(conn, only)
        if only and not projs:
            return None
        sched = _schedule(conn, only)
        todo = [e for e in _needs_setup(conn, only)
                if e in projs and ((sched.get(e) or {}).get("status") or "needs_attention") not in PAST_STATUSES]
    for e in todo:
        try:
            _setup(e)
        except Exception:
            pass   # the row still renders from whatever exists
    _preload()
    with engine.connect() as conn:
        pms = _pms(conn, only)
        est_dates = _estimate_dates(conn, only)
        sigs = _signatures(conn, only)
        rows = []
        for eid, proj in projs.items():
            comp = _billing(conn, eid, sigs.get(eid, ()))
            rows.append(_row(eid, proj, sched.get(eid), pms.get(eid, []), est_dates, comp, today))
    return rows


# ── endpoints ───────────────────────────────────────────────────────────────
@router.get("/invoicing")
def invoicing(user=Depends(get_current_user)):
    _require(user)
    rows = _build()
    return {"today": date.today().isoformat(), "soon_days": SOON_DAYS, "rows": rows}


@router.get("/project/{entity_id}")
def project_row(entity_id: str, user=Depends(get_current_user)):
    _require(user)
    rows = _build(only=entity_id)
    if not rows:
        raise HTTPException(status_code=404, detail="Project not found")
    return rows[0]


@router.get("/project/{entity_id}/history")
def project_history(entity_id: str, user=Depends(get_current_user)):
    """Billing changes on this project (invoice lines + terms + crew splits):
    who, when, what changed — newest first."""
    _require(user)
    with engine.connect() as conn:
        ms_ids = [str(r[0]) for r in conn.execute(text("""
            SELECT m.id FROM project_invoice_milestones m
            JOIN project_invoice_schedules s ON s.id = m.schedule_id WHERE s.entity_id = :e
        """), {"e": entity_id}).all()]
        q = text(f"""
            SELECT a.id, a.action, a.target_type, a.target_id, a.target_label, a.detail, a.created_at,
                   COALESCE(NULLIF(TRIM(CONCAT(COALESCE(u.first_name,''),' ',COALESCE(u.last_name,''))), ''),
                            a.actor_email) AS who
            FROM audit_log a LEFT JOIN users u ON u.id = a.actor_user_id
            WHERE (a.action LIKE 'billing.%' AND a.target_type = 'project' AND a.target_id = :e)
               OR (a.target_type = 'invoice_milestone' AND (
                     JSON_UNQUOTE(JSON_EXTRACT(a.detail, '$.entity_id')) = :e
                     {"OR a.target_id IN :ids" if ms_ids else ""}))
            ORDER BY a.id DESC LIMIT 200
        """)
        if ms_ids:
            q = q.bindparams(bindparam("ids", expanding=True))
        rows = conn.execute(q, {"e": entity_id, **({"ids": ms_ids} if ms_ids else {})}).mappings().all()
    out = []
    for r in rows:
        try:
            det = json.loads(r["detail"]) if r["detail"] else {}
        except Exception:
            det = {}
        out.append({"id": r["id"], "action": r["action"], "label": r["target_label"],
                    "who": r["who"], "at": str(r["created_at"]) if r["created_at"] else None,
                    "changes": det.get("changes"), "deleted": det.get("deleted"), "via": det.get("via")})
    return {"history": out}


@router.get("/pm-schedule")
def pm_schedule(user=Depends(get_current_user)):
    """Every active PM assignment: one entry per (PM, schedule item) with the
    item's window and the project's status — the per-PM table + timeline."""
    _require(user)
    with engine.connect() as conn:
        rows = conn.execute(text("""
            SELECT m.id AS pm_id, m.first_name, m.last_name, m.color,
                   qc.qbo_id, qc.display_name, pc.display_name AS customer, psi.id AS item_id, psi.status,
                   psi.start_date, psi.end_date,
                   DATE_ADD(psi.end_date, INTERVAL COALESCE(psi.overage_days, 0) DAY) AS true_end,
                   MAX(spm.is_primary) AS is_primary
            FROM project_schedule_item_project_managers spm
            JOIN project_schedule_items psi ON psi.id = spm.schedule_item_id
            JOIN projects p ON p.id = psi.project_id
            JOIN qbo_customers qc ON qc.id = p.qbo_customer_id
            LEFT JOIN qbo_customers pc ON pc.qbo_id = qc.parent_qbo_id
            JOIN project_managers m ON m.id = spm.project_manager_id
            WHERE spm.unassigned_at IS NULL AND m.is_active = 1
            GROUP BY m.id, m.first_name, m.last_name, m.color, qc.qbo_id, qc.display_name, pc.display_name,
                     psi.id, psi.status, psi.start_date, psi.end_date, true_end
            ORDER BY m.first_name, m.last_name, psi.start_date
        """)).mappings().all()
    pms = {}
    for r in rows:
        pm = pms.setdefault(r["pm_id"], {
            "id": r["pm_id"], "name": f"{r['first_name'] or ''} {r['last_name'] or ''}".strip(),
            "initials": _initials(r["first_name"], r["last_name"]), "color": r["color"], "items": []})
        start, end = _d(r["start_date"]), _d(r["true_end"] or r["end_date"])
        pm["items"].append({
            "entity_id": str(r["qbo_id"]), "name": r["display_name"], "customer": r["customer"],
            "item_id": r["item_id"],
            "status": r["status"] or "needs_attention", "primary": bool(r["is_primary"]),
            "start": start, "end": end,
            "days": ((date.fromisoformat(end) - date.fromisoformat(start)).days + 1) if start and end else None,
        })
    return {"today": date.today().isoformat(), "pms": sorted(pms.values(), key=lambda p: p["name"])}


# ── Progress Invoice sheet (Ashley's workbook, column for column) ───────────
# A–I mirror QuickBooks' "Estimates & Progress Invoicing Summary by Customer"
# report (open = TxnStatus 'Accepted' estimates, grouped by the project's fully
# qualified customer name, with a "Total for …" row per group). J–O are the
# columns Ashley types in from the Job Schedule; here they come from the app's
# assignment data (schedule items + PMs), the synced Job Schedule (job name)
# and the pipeline (labor days).
SHEET_HEAD = ["", "Date", "Num", "Status", "Customer", "Amount", "Invoiced amount", "% invoiced",
              "Remaining amount", "Job Name", "PM", "Invoice Dates", "Start Date", "End Date", "Job Days"]


def _open_estimates(conn):
    """Open (Accepted) estimates with QBO's progress-invoiced amount: the sum of
    the latest version of every invoice whose LinkedTxn points at the estimate."""
    ests = conn.execute(text("""
        SELECT t.qbo_id, t.doc_number, t.txn_date, t.total_amt, t.customer_qbo_id,
               qc.display_name, COALESCE(JSON_UNQUOTE(JSON_EXTRACT(qc.raw_json, '$.FullyQualifiedName')),
                                         qc.display_name) AS fqn
        FROM qbo_transactions t JOIN qbo_customers qc ON qc.qbo_id = t.customer_qbo_id
        WHERE t.entity_type = 'Estimate' AND t.total_amt > 0
          AND JSON_UNQUOTE(JSON_EXTRACT(t.raw_json, '$.TxnStatus')) = 'Accepted'
    """)).mappings().all()
    linked = conn.execute(text("""
        WITH latest AS (
          SELECT qt.total_amt, qt.raw_json,
                 ROW_NUMBER() OVER (PARTITION BY qt.qbo_id ORDER BY qt.id DESC) AS rn
          FROM qbo_transactions qt WHERE qt.entity_type = 'Invoice'
        )
        SELECT jt.txn_id, SUM(l.total_amt) AS amt
        FROM latest l, JSON_TABLE(l.raw_json, '$.LinkedTxn[*]' COLUMNS (
               txn_id VARCHAR(32) PATH '$.TxnId', txn_type VARCHAR(32) PATH '$.TxnType')) jt
        WHERE l.rn = 1 AND jt.txn_type = 'Estimate'
        GROUP BY jt.txn_id
    """)).all()
    inv = {str(r[0]): float(r[1] or 0) for r in linked}
    return [{**dict(e), "invoiced": inv.get(str(e["qbo_id"]), 0.0)} for e in ests]


def _job_names(conn):
    """Project # -> the Job Schedule's short job name ("Exxon - Baytown"), from
    the synced Job Schedule's most common description for that job."""
    try:
        rows = conn.execute(text("""
            SELECT job, descr, COUNT(*) AS n FROM gsync_schedule_days
            WHERE job IS NOT NULL AND descr IS NOT NULL GROUP BY job, descr
        """)).all()
    except Exception:
        return {}
    best = {}
    for job, descr, n in rows:
        if n > best.get(job, ("", 0))[1]:
            best[job] = (descr, n)
    return {j: " - ".join(p.strip() for p in d.split(" - ")[:2]) for j, (d, _) in best.items()}


def _labor_days(conn):
    """Quote # -> labor days from the pipeline (B151) — the Job Schedule's 'Job Days'."""
    rows = conn.execute(text("""
        SELECT quote_number, MAX(labor_days) FROM opportunities
        WHERE quote_number IS NOT NULL GROUP BY quote_number
    """)).all()
    return {str(r[0]).strip(): float(r[1]) for r in rows if r[1] is not None}


def _fallback_job_name(name):
    """'5082 - Exxon - Baytown, TX' -> 'Exxon - TX' when the Job Schedule has no entry."""
    parts = [p.strip() for p in (name or "").split(" - ")]
    if len(parts) >= 3:
        st = parts[-1].rsplit(",", 1)
        return f"{parts[1]} - {st[1].strip() if len(st) == 2 else parts[-1]}"
    return parts[-1] if parts else ""


def _num_val(v):
    return int(v) if v is not None and float(v).is_integer() else v


def _sheet(q="", pm="", status=""):
    rows = {r["entity_id"]: r for r in (_build() or [])}
    with engine.connect() as conn:
        ests = _open_estimates(conn)
        jobs = _job_names(conn)
        labor = _labor_days(conn)
    q = (q or "").strip().lower()
    groups = {}
    for e in ests:
        eid = str(e["customer_qbo_id"])
        r = rows.get(eid) or {}
        if pm and not any(str(p["id"]) == str(pm) for p in r.get("pms", [])):
            continue
        if status and (r.get("status") or "needs_attention") != status:
            continue
        name = e["display_name"] or ""
        number = name.split(" - ")[0].strip() if " - " in name else ""
        doc = (e["doc_number"] or "").strip()
        job = jobs.get(number) or _fallback_job_name(name)
        if q and q not in f"{e['fqn']} {name} {doc} {job}".lower():
            continue
        # L: this estimate's invoice dates still to send (Billing & Schedule plan)
        est = next((x for x in r.get("estimates", []) if str(x["qbo_id"]) == str(e["qbo_id"])), None)
        inv_dates = sorted({m["invoice_date"] for m in (est or {}).get("milestones", [])
                            if m["tier"] in ("scheduled", "partial") and m.get("invoice_date")})
        # O: the estimate's own labor days; a change order without its own → N/A
        days = labor.get(doc)
        if not days:
            base = doc.split(" - ")[0].strip()      # "4456 - 1" is still job 4456
            days = "N/A" if doc and base != number else (labor.get(number) or "?")
        amt = round(float(e["total_amt"] or 0), 2)
        invd = round(e["invoiced"], 2)
        pms = r.get("pms", [])
        groups.setdefault(e["fqn"], []).append({
            "entity_id": eid, "estimate_qbo_id": str(e["qbo_id"]),
            "date": e["txn_date"].strftime("%m/%d/%Y") if e["txn_date"] else "",
            "_sort": str(e["txn_date"] or ""),
            "num": doc, "status": "Accepted", "customer": name,
            "amount": amt, "invoiced": invd or None,
            "pct": round(invd / amt, 4) if amt else 0, "remaining": round(amt - invd, 2),
            "job_name": job,
            "pm": "/".join(p["initials"] for p in pms) or "?", "pm_ids": [str(p["id"]) for p in pms],
            "invoice_dates": inv_dates, "start": r.get("start"), "end": r.get("end"),
            "job_days": _num_val(days) if not isinstance(days, str) else days,
            "sched_status": r.get("status") or "needs_attention",
        })
    out = []
    for fqn in sorted(groups, key=lambda s: s.lower()):
        items = sorted(groups[fqn], key=lambda x: (x["_sort"], x["num"]), reverse=True)
        for x in items:
            x.pop("_sort", None)
        out.append({"name": fqn, "rows": items,
                    "amount": round(sum(x["amount"] for x in items), 2),
                    "invoiced": round(sum(x["invoiced"] or 0 for x in items), 2)})
    return out


def _md(iso):
    y, m, d = iso.split("-")
    return f"{int(m)}/{int(d)}"


def _xlsx(groups):
    """Build the workbook the way Ashley's file looks: QBO report styling on
    A–I (Arial 8, bold group/total rows, $ totals with a top rule), Calibri 11
    centered on J–O, header rule, frozen header + filter; plus Sheet2, the
    per-PM job list."""
    import io
    from datetime import datetime
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, Side

    thin = Side(style="thin")
    a8, a8b = Font(name="Arial", size=8), Font(name="Arial", size=8, bold=True)
    c11, c11b = Font(name="Calibri", size=11), Font(name="Calibri", size=11, bold=True)
    wrap, center = Alignment(wrap_text=True), Alignment(horizontal="center", vertical="center")

    def dt(iso):
        return datetime.strptime(iso, "%Y-%m-%d") if iso else "?"

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for col, w in zip("ABCDEFGHIJK", (85.75, 9.25, 7.5, 8.375, 47.0, 16.125, None, 9.25, 18.625, 21.25, 11.25)):
        if w:
            ws.column_dimensions[col].width = w
    ws.append(SHEET_HEAD)
    for i in range(2, 16):
        c = ws.cell(1, i)
        c.border = Border(bottom=thin)
        if i <= 9:
            c.font, c.alignment = Font(name="Arial", size=9, bold=True), Alignment(horizontal="center", wrap_text=True)
        else:
            c.font, c.alignment = c11b, Alignment(horizontal="center")
    r = 1
    for g in groups:
        r += 1
        c = ws.cell(r, 1, g["name"])
        c.font, c.alignment = a8, Alignment(horizontal="left", wrap_text=True)
        top = r + 1
        for x in g["rows"]:
            r += 1
            ds = x["invoice_dates"]
            vals = [None, x["date"], x["num"], x["status"], x["customer"], x["amount"], x["invoiced"], x["pct"],
                    x["remaining"], x["job_name"], x["pm"],
                    (dt(ds[0]) if len(ds) == 1 else ", ".join(_md(d) for d in ds) if ds else "?"),
                    dt(x["start"]), dt(x["end"]), x["job_days"]]
            for i, v in enumerate(vals, start=1):
                c = ws.cell(r, i, v)
                if 2 <= i <= 9:
                    c.font, c.alignment = a8, wrap
                    if i in (6, 7, 9):
                        c.number_format = "#,##0.00"
                    elif i == 8:
                        c.number_format = "#,##0.00%"
                elif i >= 10:
                    c.font = c11
                    if i > 10:
                        c.alignment = center
                    if isinstance(v, datetime):
                        c.number_format = "mm-dd-yy"
        r += 1
        c = ws.cell(r, 1, f"Total for {g['name']}")
        c.font, c.alignment = a8b, Alignment(horizontal="left", wrap_text=True)
        for col in ("F", "G"):
            c = ws[f"{col}{r}"]
            c.value = f"=SUM({col}{top}:{col}{r - 1})"
            c.font, c.alignment, c.number_format = a8b, wrap, '"$"#,##0.00'
            c.border = Border(top=thin)
    if groups:
        r += 1
        ws.cell(r, 1, "TOTAL").font = a8b
        for col in ("F", "G"):
            c = ws[f"{col}{r}"]
            c.value = f'=SUMIF(A2:A{r - 1},"Total for*",{col}2:{col}{r - 1})'
            c.font, c.number_format = a8b, '"$"#,##0.00'
            c.border = Border(top=thin, bottom=Side(style="double"))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:O{max(r, 1)}"

    # Sheet2: one block per PM — Job Name, PM, Start, End, Job Days + a count.
    ws2 = wb.create_sheet("Sheet2")
    for col, w in (("A", 21.25), ("B", 3.5), ("C", 9.625), ("E", 7.5)):
        ws2.column_dimensions[col].width = w
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    jobs = {}
    for g in groups:
        for x in g["rows"]:
            j = jobs.setdefault((x["pm"], x["job_name"]), x)
            if not isinstance(j["job_days"], (int, float)) and isinstance(x["job_days"], (int, float)):
                jobs[(x["pm"], x["job_name"])] = x
    r = 0
    for pm_i in sorted({k[0] for k in jobs}, key=lambda i: (i == "?", i)):   # unassigned last
        items = [v for k, v in jobs.items() if k[0] == pm_i]
        r += 1
        for i, h in enumerate(SHEET_HEAD[9:10] + SHEET_HEAD[10:11] + SHEET_HEAD[12:15], start=1):
            c = ws2.cell(r, i, h)
            c.font, c.alignment, c.border = c11b, Alignment(horizontal="center"), box
        for x in items:
            r += 1
            for i, v in enumerate([x["job_name"], x["pm"], dt(x["start"]), dt(x["end"]), x["job_days"]], start=1):
                c = ws2.cell(r, i, v)
                c.font, c.border = c11, box
                if i > 1:
                    c.alignment = center
                if isinstance(v, datetime):
                    c.number_format = "mm-dd-yy"
        r += 1
        ws2.cell(r, 1, len(items))
        r += 2
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _csv(groups):
    import csv
    import io
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(SHEET_HEAD)
    for g in groups:
        w.writerow([g["name"]] + [""] * 14)
        for x in g["rows"]:
            ds = x["invoice_dates"]
            w.writerow(["", x["date"], x["num"], x["status"], x["customer"], f"{x['amount']:.2f}",
                        f"{x['invoiced']:.2f}" if x["invoiced"] else "", f"{x['pct'] * 100:.2f}%",
                        f"{x['remaining']:.2f}", x["job_name"], x["pm"],
                        ", ".join(_md(d) for d in ds) if ds else "?",
                        x["start"] or "?", x["end"] or "?", x["job_days"]])
        w.writerow([f"Total for {g['name']}", "", "", "", "", f"{g['amount']:.2f}", f"{g['invoiced']:.2f}"] + [""] * 8)
    return buf.getvalue()


@router.get("/progress-sheet")
def progress_sheet(q: str = "", pm: str = "", status: str = "", user=Depends(get_current_user)):
    _require(user)
    return {"today": date.today().isoformat(), "columns": SHEET_HEAD, "groups": _sheet(q, pm, status)}


@router.get("/progress-sheet/export")
def progress_sheet_export(fmt: str = "xlsx", q: str = "", pm: str = "", status: str = "",
                          user=Depends(get_current_user)):
    _require(user)
    from fastapi.responses import Response
    groups = _sheet(q, pm, status)
    stamp = date.today().strftime("%m-%d-%y")
    if fmt == "csv":
        return Response(content=_csv(groups).encode("utf-8-sig"), media_type="text/csv",
                        headers={"Content-Disposition": f'attachment; filename="Progress Invoice {stamp}.csv"'})
    return Response(content=_xlsx(groups),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="Progress Invoice {stamp}.xlsx"'})


def _warm():
    """Build the per-project billing cache once at startup so the first visit
    after a deploy doesn't wait for every project to compose."""
    import time
    time.sleep(20)
    try:
        _build()
    except Exception:
        pass


threading.Thread(target=_warm, daemon=True).start()
