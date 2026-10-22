"""
Read OPI's two working Google Sheets from Drive (read-only) and keep the latest
parsed copy in the database:

  * "Updated OP.B151.Project Forecast" → gsync_b151_rows (tab "2. Rolling
    Revenue": header row 10, one row per quote);
  * "1 ON POINT Job Schedule.xlsx"     → gsync_schedule_days (two-month tabs,
    weekly blocks of crew rows; each day is 5 columns Crw | Job# | Description |
    extra | Date, a job repeats its number in every day cell it covers).

Nothing here writes to Drive.
"""
import datetime as dt
import io
import json
import re
import warnings

from sqlalchemy import text

from app import gdrive

B151_ID = "1bw6sWq94xL2Hfp4jE99QRjsPRahnXWgZO6m3DgK5wgg"
JOB_SCHEDULE_ID = "1W-KZD4Yn7S2AAnUj_7d0FI9rq1J2XV3i"

# "2. Rolling Revenue" columns (0-based), from header row 10.
RR = {
    "status": 1, "last_contact": 2, "follow_ups": 3, "comm_type": 4, "link": 5, "quote": 6,
    "requested": 7, "revised": 8, "revisions": 9, "order_date": 10, "quoted_by": 11,
    "contact": 12, "customer": 13, "city": 14, "state": 15, "description": 16,
    "labor_days": 17, "travel_days": 18, "start": 19, "end": 20,
    "ohp_amount": 25, "ohp_pct": 26, "tcv": 27, "sent": 28, "order_value": 29,
    "decision": 31, "discounted": 32,
}
RR_HEADER_ROW = 10
_MON = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}


def _openpyxl():
    import openpyxl  # imported lazily: only the sync needs it
    warnings.filterwarnings("ignore", module="openpyxl")
    return openpyxl


def _val(v):
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def _quote(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    s = str(v).strip()
    return s or None


# ── B151 ─────────────────────────────────────────────────────────────────────

DATE_COLS = {"last_contact", "requested", "revised", "order_date", "start", "end", "sent", "decision"}


def _serial_date(v):
    if isinstance(v, (int, float)) and 20000 < v < 80000:
        return (dt.date(1899, 12, 30) + dt.timedelta(days=int(v))).isoformat()
    return _val(v)


def _rows_from_values(values):
    out = []
    for i, row in enumerate(values, start=RR_HEADER_ROW + 1):
        row = list(row) + [None] * (40 - len(row))
        q = _quote(row[RR["quote"]])
        if not q:
            continue
        rec = {k: (_serial_date(row[c]) if k in DATE_COLS else _val(row[c])) for k, c in RR.items()}
        rec["quote"] = q
        rec["row"] = i
        out.append(rec)
    return out


def fetch_b151_rows(min_expected=0):
    """The Rolling Revenue rows. Sheets API first (exact cell values); if it's
    not enabled, the xlsx export — which Google sometimes returns with rows
    missing on a big sheet, so it's retried and the fullest copy kept."""
    meta = gdrive.get(B151_ID)
    try:
        vals = gdrive.sheet_values(B151_ID, f"'2. Rolling Revenue'!A{RR_HEADER_ROW + 1}:AH6000")
        rows = _rows_from_values(vals)
        try:
            _attach_links(rows, gdrive.sheet_links(B151_ID, f"'2. Rolling Revenue'!F{RR_HEADER_ROW + 1}:F6000",
                                                   RR_HEADER_ROW + 1))
        except Exception:
            pass
        return meta, rows, "sheets_api"
    except Exception:
        pass
    best, best_bytes = [], None
    for _ in range(4):
        data = gdrive.export(B151_ID, gdrive.XLSX)
        rows = parse_b151(data)
        if len(rows) > len(best):
            best, best_bytes = rows, data
        if min_expected and len(best) >= min_expected:
            break
    if best_bytes:
        _attach_links(best, xlsx_hyperlinks(best_bytes, "2. Rolling Revenue", "F"))
    return meta, best, "export"


def _attach_links(rows, links):
    for r in rows:
        if links.get(r["row"]):
            r["link_url"] = links[r["row"]]


def xlsx_hyperlinks(xlsx_bytes, sheet_name, col):
    """{row: url} for external hyperlinks in one column, read straight from the
    xlsx XML (cheap — no full workbook load)."""
    import zipfile
    z = zipfile.ZipFile(io.BytesIO(xlsx_bytes))
    wb = z.read("xl/workbook.xml").decode("utf-8")
    m = re.search(r'<sheet[^>]*name="%s"[^>]*r:id="(rId\d+)"' % re.escape(sheet_name), wb)
    if not m:
        return {}
    rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    t = re.search(r'<Relationship[^>]*Id="%s"[^>]*/>' % m.group(1), rels)
    target = re.search(r'Target="([^"]+)"', t.group(0)).group(1).lstrip("/")
    path = target if target.startswith("xl/") else "xl/" + target
    sheet_xml = z.read(path).decode("utf-8")
    rel_path = "xl/worksheets/_rels/" + path.split("/")[-1] + ".rels"
    if rel_path not in z.namelist():
        return {}
    targets = dict(re.findall(r'Id="(rId\d+)"[^>]*Target="([^"]+)"', z.read(rel_path).decode("utf-8")))
    targets.update({k: v for v, k in re.findall(r'Target="([^"]+)"[^>]*Id="(rId\d+)"', z.read(rel_path).decode("utf-8"))})
    out = {}
    for h in re.findall(r"<hyperlink [^>]*>", sheet_xml):
        ref = re.search(r'ref="([A-Z]+)(\d+)"', h)
        rid = re.search(r'r:id="(rId\d+)"', h)
        if ref and rid and ref.group(1) == col and targets.get(rid.group(1), "").startswith("http"):
            out[int(ref.group(2))] = targets[rid.group(1)].replace("&amp;", "&")
    return out


def parse_b151(xlsx_bytes):
    """Rows of "2. Rolling Revenue" that carry a quote # → list of dicts."""
    op = _openpyxl()
    wb = op.load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    ws = wb["2. Rolling Revenue"]
    out = []
    for i, row in enumerate(ws.iter_rows(min_row=RR_HEADER_ROW + 1, max_col=40, values_only=True),
                            start=RR_HEADER_ROW + 1):
        row = list(row) + [None] * (40 - len(row))
        q = _quote(row[RR["quote"]])
        if not q:
            continue
        rec = {k: _val(row[c]) for k, c in RR.items()}
        rec["quote"] = q
        rec["row"] = i
        out.append(rec)
    wb.close()
    return out


def store_b151(conn, rows):
    conn.execute(text("DELETE FROM gsync_b151_rows"))
    if rows:
        conn.execute(text("INSERT INTO gsync_b151_rows (row_num, quote_number, data) VALUES (:n, :q, :d)"),
                     [{"n": r["row"], "q": r["quote"], "d": json.dumps(r, default=str)} for r in rows])


def load_b151(conn):
    return [json.loads(d) if isinstance(d, str) else d
            for d in conn.execute(text("SELECT data FROM gsync_b151_rows ORDER BY row_num")).scalars()]


# ── Job Schedule ─────────────────────────────────────────────────────────────

def fetch_job_schedule():
    meta = gdrive.get(JOB_SCHEDULE_ID)
    data = gdrive.download(JOB_SCHEDULE_ID)
    return meta, data


def _tab_months(name):
    """'Sept-Oct 2026' → (2026, [9, 10]); None for non-calendar tabs."""
    m = re.search(r"(20\d\d)", name)
    if not m:
        return None
    mons = [_MON[t[:3].lower()] for t in re.findall(r"[A-Za-z]{3,}", name) if t[:3].lower() in _MON]
    return (int(m.group(1)), mons) if mons else None


def _fill(c):
    f = getattr(c, "fill", None)
    if not f or not f.fill_type:
        return None
    fg = f.fgColor
    if fg.type == "rgb":
        return fg.rgb
    if fg.type == "theme":
        return f"t{fg.theme}"
    return f"i{fg.indexed}"


def _font(c):
    """Font colour of a cell (OPI types overage days in orange text)."""
    col = getattr(getattr(c, "font", None), "color", None)
    if col is None:
        return None
    try:
        if col.type == "rgb":
            return col.rgb if isinstance(col.rgb, str) else None
        return f"t{col.theme}" if col.type == "theme" else f"i{col.indexed}"
    except Exception:
        return None


def parse_job_schedule(xlsx_bytes, since=None):
    """Day cells for calendar tabs that reach `since` or later."""
    op = _openpyxl()
    wb = op.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)
    recs = {}
    for name in wb.sheetnames:
        tm = _tab_months(name)
        if not tm:
            continue
        year, mons = tm
        last = dt.date(year + (1 if mons[-1] < mons[0] else 0), mons[-1], 28)
        if since and last < since.replace(day=1):
            continue
        ws = wb[name]
        dates = None
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=35):
            vals = [c.value for c in row] + [None] * (35 - len(row))
            dv = [vals[5 * d + 4] for d in range(7)]
            if sum(isinstance(x, dt.datetime) for x in dv) >= 4:
                dates = []
                for d in range(7):
                    x = dv[d]
                    if isinstance(x, dt.datetime):
                        y = year
                        if mons[0] == 1 and x.month == 12:
                            y -= 1
                        elif mons[0] >= 11 and x.month <= 2:
                            y += 1
                        dates.append(dt.date(y, x.month, x.day))
                    else:
                        dates.append(None)
                for d in range(7):
                    if dates[d] is None:
                        for e in range(7):
                            if dates[e]:
                                dates[d] = dates[e] + dt.timedelta(days=d - e)
                                break
                continue
            if not dates:
                continue
            for d in range(7):
                crew = str(vals[5 * d] or "").strip().upper()
                if not re.fullmatch(r"[A-Z]{2,4}\d?", crew):
                    continue
                job, desc, ex = (_val(vals[5 * d + k]) for k in (1, 2, 3))
                job = str(job) if job is not None else None
                desc = str(desc) if desc is not None else None
                ex = str(ex) if ex is not None else None
                if not (job or desc or ex) or not dates[d]:
                    continue
                day = dates[d]
                own_month = day.month in mons
                key = (day, crew)
                if key in recs and recs[key]["own"] and not own_month:
                    continue   # the tab that owns the month wins over a spill-over week
                recs[key] = {"tab": name, "day": day, "crew": crew, "job": job, "descr": desc, "extra": ex,
                             "job_fill": _fill(row[5 * d + 1]) if len(row) > 5 * d + 1 else None,
                             "desc_fill": _fill(row[5 * d + 2]) if len(row) > 5 * d + 2 else None,
                             "desc_font": _font(row[5 * d + 2]) if len(row) > 5 * d + 2 else None,
                             "own": own_month}
    wb.close()
    return list(recs.values())


def store_job_schedule(conn, recs):
    conn.execute(text("DELETE FROM gsync_schedule_days"))
    if recs:
        conn.execute(text("""
            INSERT INTO gsync_schedule_days (tab, day, crew, job, descr, extra, job_fill, desc_fill, desc_font)
            VALUES (:tab, :day, :crew, :job, :descr, :extra, :job_fill, :desc_fill, :desc_font)
        """), [{**r, "job": (r["job"] or "")[:120] or None, "descr": (r["descr"] or "")[:500] or None,
                "extra": (r["extra"] or "")[:500] or None, "desc_font": r.get("desc_font")} for r in recs])
