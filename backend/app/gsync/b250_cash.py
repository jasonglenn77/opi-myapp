"""
OPI's "B250.OP.13 Week Cash Flow Forecast" workbook → gsync_b250_snapshot /
gsync_b250_lines, so Data Review (review_sources.py) can check each typed number
against QuickBooks and the crew Payment Schedules. Nothing in the forecast is
taken from here; it is a copy to compare against.

The workbook: one tab per week ("October 26-1"), copied forward by hand each
Monday; the current tab is the one whose D6 says "Ending-->" with the latest
week-1 date in E6. Columns E..Q = weeks 1..13 (week ending Fridays). Also the
"SnapShot" (past dues + collection notes) and "MonthlyYearly Totals" tabs.

Source: Drive (env B250_FILE_ID, or a file named like "B250 … Cash Flow" shared
with the service account), else an xlsx uploaded on the Data Review page.

  python -m app.gsync.b250_cash load <file.xlsx>    # load a local copy (source 'upload')
  python -m app.gsync.b250_cash drive               # pull from Drive
"""
import io
import os
import re
import sys
from datetime import date, datetime

from sqlalchemy import bindparam, text

from app.db import engine

KEEP = 5
WEEK_COLS = range(5, 18)            # E..Q
GSHEET = "application/vnd.google-apps.spreadsheet"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MONTHS = ["january", "february", "march", "april", "may", "june", "july", "august", "september",
          "october", "november", "december"]


class NotAccessible(RuntimeError):
    pass


def _d(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def _num(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("$", "").replace(",", "")
    try:
        return float(s)
    except ValueError:
        return None


def _s(v):
    return str(v).strip() if v is not None and str(v).strip() else None


def _col(c):
    from openpyxl.utils import get_column_letter
    return get_column_letter(c)


def _current_tab(wb):
    best = None
    for ws in wb.worksheets:
        if str(ws.cell(6, 4).value or "").strip().lower().startswith("ending"):
            d = _d(ws.cell(6, 5).value)
            if d and (best is None or d > best[0] or (d == best[0] and ws.sheet_state == "visible")):
                best = (d, ws.title)
    return best


def parse(content):
    """xlsx bytes → (tab_name, week1_end, lines)."""
    import openpyxl
    wf = openpyxl.load_workbook(io.BytesIO(content), data_only=False)
    wv = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    cur = _current_tab(wf)
    if not cur:
        raise ValueError("No weekly tab found (D6 'Ending-->' with a date in E6) — is this the B250 workbook?")
    week1, tab = cur
    f, v = wf[tab], wv[tab]
    weeks = {}
    for c in WEEK_COLS:
        d = _d(v.cell(6, c).value)
        if d:
            weeks[c] = d
    lines = []

    def add(section, r, c=None, **kw):
        cell_f = f.cell(r, c).value if c else None
        isf = isinstance(cell_f, str) and cell_f.startswith("=")
        amt = _num(v.cell(r, c).value) if c else kw.pop("amount", None)
        lines.append({"section": section, "row_num": r, "col": _col(c) if c else None,
                      "week_end": weeks.get(c) if c else None, "amount": amt,
                      "is_formula": 1 if isf else 0, "formula": cell_f[:500] if isf else None, **kw})

    section, dealer = None, None
    for r in range(9, f.max_row + 1):
        b = _s(f.cell(r, 2).value)
        if not b:
            continue
        low = b.lower()
        if low.startswith("opening"):
            add("opening", r, 5, label=b)
            for c in WEEK_COLS:
                if c != 5:
                    add("total", r, c, label=b)
            continue
        if low.startswith("cash inflow"):
            section = "inflow"
            for c in WEEK_COLS:
                add("total", r, c, label=b)
            continue
        if low.startswith("cash outflow"):
            section = "outflow"
            for c in WEEK_COLS:
                add("total", r, c, label=b)
            continue
        if low.startswith("total cash") or low.startswith("ending cash"):
            for c in WEEK_COLS:
                add("total", r, c, label=b)
            if low.startswith("ending cash"):
                break
            continue
        if section == "inflow":
            inv = _s(f.cell(r, 3).value)
            if low.startswith("forecasted invoice"):
                sec, kw = "forecast_invoice", {"label": b, "group_label": dealer}
            elif inv:
                sec, kw = "inflow_invoice", {"label": b, "group_label": dealer,
                                             "invoice_num": inv.split(".")[0][:40], "due_date": _d(f.cell(r, 4).value)}
            else:
                dealer = b
                sec, kw = "inflow_dealer", {"label": b, "group_label": b}
            vals = [c for c in WEEK_COLS if _num(v.cell(r, c).value) not in (None, 0.0)]
            if sec == "inflow_invoice" and not vals:
                add(sec, r, None, **kw)
            for c in vals:
                add(sec, r, c, **dict(kw))
            continue
        if section in ("outflow", "ap", "contractor", "payroll", "operating"):
            if low == "a/p":
                section = "ap"
                for c in WEEK_COLS:
                    add("total", r, c, label=b)
                continue
            if low == "contractors":
                section = "contractor"
                continue
            if low == "payroll":
                section = "payroll"
                continue
            if low.startswith("other operating"):
                section = "operating"
                continue
            for c in WEEK_COLS:
                val = _num(v.cell(r, c).value)
                if val not in (None, 0.0) or isinstance(f.cell(r, c).value, str):
                    add(section if section != "outflow" else "operating", r, c, label=b)

    if "SnapShot" in wf.sheetnames:
        s = wv["SnapShot"]
        for r in range(2, s.max_row + 1):
            job = _s(s.cell(r, 1).value)
            if not job:
                continue
            lines.append({"section": "snapshot", "row_num": r, "col": "I", "label": job,
                          "group_label": _s(s.cell(r, 2).value), "due_date": _d(s.cell(r, 8).value),
                          "week_end": _d(s.cell(r, 7).value), "amount": _num(s.cell(r, 9).value),
                          "note": (" | ".join(x for x in (f"past due {s.cell(r, 10).value}" if s.cell(r, 10).value is not None else None,
                                                          _s(s.cell(r, 12).value)) if x) or None),
                          "is_formula": 0, "formula": None})
    if "MonthlyYearly Totals" in wf.sheetnames:
        m = wv["MonthlyYearly Totals"]
        for (yc, bc, rc) in ((1, 2, 3), (5, 6, 7)):
            yr = _num(m.cell(1, yc).value)
            if not yr:
                continue
            for r in range(3, 15):
                name = str(m.cell(r, yc).value or "").strip().lower()
                if name not in MONTHS:
                    continue
                first = date(int(yr), MONTHS.index(name) + 1, 1)
                for lab, c in (("bookings", bc), ("revenue", rc)):
                    val = _num(m.cell(r, c).value)
                    if val is not None:
                        lines.append({"section": "monthly", "row_num": r, "col": _col(c), "label": lab,
                                      "week_end": first, "amount": val, "is_formula": 0, "formula": None})
    return tab, week1, lines


def store(content, source, who, file_id=None, file_name=None, file_modified=None):
    tab, week1, lines = parse(content)
    cols = ("section", "row_num", "col", "label", "group_label", "crew_code", "invoice_num", "due_date", "week_end",
            "amount", "is_formula", "formula", "note")
    with engine.begin() as conn:
        sid = conn.execute(text("""
            INSERT INTO gsync_b250_snapshot (source, file_id, file_name, file_modified, tab_name, week1_end, line_count, loaded_by)
            VALUES (:s, :fid, :fn, :fm, :t, :w, :n, :by)"""),
            {"s": source, "fid": file_id, "fn": (file_name or "")[:255] or None, "fm": file_modified, "t": tab,
             "w": week1, "n": len(lines), "by": who}).lastrowid
        rows = [{**{k: ln.get(k) for k in cols}, "sid": sid} for ln in lines]
        for r in rows:
            for k in ("label", "group_label"):
                if r[k]:
                    r[k] = r[k][:255]
            if r["note"]:
                r["note"] = r["note"][:1000]
        if rows:
            conn.execute(text(f"""INSERT INTO gsync_b250_lines (snapshot_id, {", ".join(cols)})
                VALUES (:sid, {", ".join(":" + c for c in cols)})"""), rows)
        old = [r[0] for r in conn.execute(text(
            "SELECT id FROM gsync_b250_snapshot ORDER BY id DESC LIMIT 1000 OFFSET :k"), {"k": KEEP})]
        if old:
            conn.execute(text("DELETE FROM gsync_b250_lines WHERE snapshot_id IN :ids").bindparams(
                bindparam("ids", expanding=True)), {"ids": old})
            conn.execute(text("DELETE FROM gsync_b250_snapshot WHERE id IN :ids").bindparams(
                bindparam("ids", expanding=True)), {"ids": old})
    return {"snapshot_id": sid, "tab": tab, "week1_end": str(week1), "lines": len(lines), "source": source}


def find_drive_file():
    from app import gdrive
    fid = os.getenv("B250_FILE_ID", "").strip()
    if fid:
        try:
            return gdrive.get(fid)
        except Exception as e:
            raise NotAccessible(f"B250_FILE_ID {fid} is not readable by the service account ({str(e)[:120]})")
    files = gdrive.list_files("name contains 'B250' and trashed = false")
    files = [x for x in files if "cash" in x["name"].lower() and x["mimeType"] in (GSHEET, XLSX)]
    if not files:
        raise NotAccessible("The B250 cash-flow workbook isn't shared with the app's Google account "
                            f"({gdrive.service_email()}) — share it (Viewer) or upload a copy on Data Review.")
    return max(files, key=lambda x: x.get("modifiedTime") or "")


def refresh_from_drive(who="cron"):
    """Pull the B250 from Drive when it's shared; skips (no new snapshot) when the
    latest snapshot already has this Drive version."""
    from app import gdrive
    f = find_drive_file()
    with engine.connect() as conn:
        last = conn.execute(text("SELECT file_id, file_modified FROM gsync_b250_snapshot ORDER BY id DESC LIMIT 1")).first()
    if last and last[0] == f["id"] and last[1] == f.get("modifiedTime"):
        return {"skipped": "unchanged", "file": f["name"]}
    content = gdrive.export(f["id"], XLSX) if f["mimeType"] == GSHEET else gdrive.download(f["id"])
    return store(content, "drive", who, f["id"], f["name"], f.get("modifiedTime"))


def latest_meta():
    try:
        with engine.connect() as conn:
            r = conn.execute(text("""SELECT id, source, file_name, tab_name, week1_end, line_count, loaded_at, loaded_by
                FROM gsync_b250_snapshot ORDER BY id DESC LIMIT 1""")).mappings().first()
    except Exception:
        return None
    return {k: (str(v) if isinstance(v, (date, datetime)) else v) for k, v in dict(r).items()} if r else None


if __name__ == "__main__":
    import json
    if len(sys.argv) > 2 and sys.argv[1] == "load":
        with open(sys.argv[2], "rb") as fh:
            print(json.dumps(store(fh.read(), "upload", "cli", file_name=os.path.basename(sys.argv[2])), default=str))
    elif len(sys.argv) > 1 and sys.argv[1] == "drive":
        try:
            print(json.dumps(refresh_from_drive("cli"), default=str))
        except NotAccessible as e:
            print("not accessible:", e)
    else:
        print(__doc__)
